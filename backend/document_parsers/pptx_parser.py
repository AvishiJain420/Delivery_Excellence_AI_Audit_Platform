import io
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt
from document_parsers.base import BaseParser
from agent.document import Document


class PptxParser(BaseParser):

    def parse(self, content_bytes: bytes, filename: str, metadata: dict) -> Document:
        try:
            presentation = Presentation(io.BytesIO(content_bytes))

            text_parts = []
            tables = []
            images = []
            content_sequence = []

            for slide_num, slide in enumerate(presentation.slides, start=1):

                # Slide header in sequence — gives context to everything below
                content_sequence.append({
                    "type": "text",
                    "value": f"\n--- Slide {slide_num} ---"
                })

                # Sort shapes by vertical position so reading order is preserved
                sorted_shapes = sorted(
                    slide.shapes,
                    key=lambda s: (s.top or 0, s.left or 0)
                )

                for shape in sorted_shapes:

                    # ── Text ─────────────────────────────────────────
                    if shape.has_text_frame:
                        shape_text = []
                        for para in shape.text_frame.paragraphs:
                            line = "".join(run.text for run in para.runs)
                            if line.strip():
                                shape_text.append(line)

                        if shape_text:
                            combined = "\n".join(shape_text)
                            text_parts.append(combined)
                            content_sequence.append({
                                "type": "text",
                                "value": combined
                            })

                    # ── Tables ───────────────────────────────────────
                    elif shape.has_table:
                        rows = [
                            [cell.text.strip() for cell in row.cells]
                            for row in shape.table.rows
                        ]
                        tables.append(rows)
                        content_sequence.append({
                            "type": "table",
                            "value": rows,
                            "label": f"Table on slide {slide_num}"
                        })

                    # ── Images ───────────────────────────────────────
                    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        try:
                            img_bytes = shape.image.blob
                            images.append(img_bytes)
                            content_sequence.append({
                                "type": "image",
                                "value": img_bytes,
                                "label": f"Slide {slide_num} image"
                            })
                        except Exception:
                            pass

            document = Document(
                id=metadata.get("item_id", filename),
                filename=filename,
                file_type="pptx",
                text="\n".join(text_parts),
                tables=tables,
                images=images,
                content_sequence=content_sequence,
                page_count=len(presentation.slides),
                metadata=metadata,
            )
            return document
         
        except Exception as e:

            print(f"  PPTX parse error for {filename}: {e}")
            document = Document(
                id=metadata.get("item_id", filename),
                filename=filename,
                file_type="pptx",
                metadata=metadata,
                parse_error=str(e),
            )
            return document